"""Convert a preserved Fleet board XLSX export into a review-only JSON snapshot.

Requires openpyxl. This script never contacts MSERP or changes the workbook.
Keep the original XLSX beside the output: it preserves formatting, formulas,
notes, and any original precision loss that cannot be recovered from cell values.
"""
import argparse
import datetime
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def cell_text(cell):
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (float, int)):
        # Explicit zero-padded unit formats are significant. Never strip them.
        if re.fullmatch(r"0+", cell.number_format or "") and float(value).is_integer():
            return str(int(value)).zfill(len(cell.number_format))
        return str(int(value)) if float(value).is_integer() else str(value)
    return str(value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--spreadsheet-id", required=True)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    snapshot = {
        "spreadsheetId": args.spreadsheet_id,
        "capturedAt": datetime.datetime.now(datetime.timezone.utc).isoformat().replace("+00:00", "Z"),
        "sheets": [],
    }
    for sheet in workbook.worksheets:
        rows = []
        for cells in sheet.iter_rows():
            values = [cell_text(cell) for cell in cells]
            while values and values[-1] == "":
                values.pop()
            rows.append(values)
        while rows and not rows[-1]:
            rows.pop()
        snapshot["sheets"].append({"name": sheet.title, "rows": rows})
    workbook.close()
    args.output.parent.mkdir(parents=True, exist_ok=True)
    # Refuse to overwrite a dated source snapshot accidentally.
    with args.output.open("x", encoding="utf-8") as output:
        json.dump(snapshot, output, ensure_ascii=False, separators=(",", ":"))
    print(f"Prepared {len(snapshot['sheets'])} tabs for review: {args.output}")


if __name__ == "__main__":
    main()
